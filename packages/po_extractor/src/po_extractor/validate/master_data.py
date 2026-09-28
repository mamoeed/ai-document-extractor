"""Load and cache the customer and item master Excel files."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, time
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable, Optional

import openpyxl

from ..errors import MasterDataError
from .normalize import normalize_id, normalize_name

CUSTOMER_SHEET = "Customer Master"
ITEM_SHEET = "Item Master"
CUSTOMER_REQUIRED = ("Customer no.", "Legal name", "Active")
ITEM_REQUIRED = ("Source item no.", "Description", "Base UoM", "Net price EUR", "Active")
HEADER_SCAN_ROWS = 10

Record = dict[str, Any]


@dataclass(frozen=True)
class CustomerMaster:
    path: str
    sheet: str
    records: tuple[Record, ...]
    _by_number: dict[str, list[Record]] = field(repr=False)
    _by_name: dict[str, list[Record]] = field(repr=False)

    def by_number(self, number: Optional[str]) -> list[Record]:
        return list(self._by_number.get(normalize_id(number) or "", []))

    def by_name(self, name: Optional[str]) -> list[Record]:
        return list(self._by_name.get(normalize_name(name) or "", []))


@dataclass(frozen=True)
class ItemMaster:
    path: str
    sheet: str
    records: tuple[Record, ...]
    _by_item_no: dict[str, list[Record]] = field(repr=False)

    def by_item_no(self, item_no: Optional[str]) -> list[Record]:
        return list(self._by_item_no.get(normalize_id(item_no) or "", []))


def load_customer_master(path: str | Path) -> CustomerMaster:
    resolved, mtime = _stat(path, "Customer master")
    return _load_customer_master(resolved, mtime)


def load_item_master(path: str | Path) -> ItemMaster:
    resolved, mtime = _stat(path, "Item master")
    return _load_item_master(resolved, mtime)


# Cache keyed by (path, mtime): an updated file is re-read automatically.
@lru_cache(maxsize=8)
def _load_customer_master(path: str, mtime_ns: int) -> CustomerMaster:
    sheet, records = _read_table(path, CUSTOMER_SHEET, CUSTOMER_REQUIRED, id_columns=("Customer no.",))
    records = [r for r in records if r.get("Customer no.")]
    return CustomerMaster(
        path=path,
        sheet=sheet,
        records=tuple(records),
        _by_number=_index(records, lambda r: normalize_id(r.get("Customer no."))),
        _by_name=_index(records, lambda r: normalize_name(r.get("Legal name"))),
    )


@lru_cache(maxsize=8)
def _load_item_master(path: str, mtime_ns: int) -> ItemMaster:
    sheet, records = _read_table(path, ITEM_SHEET, ITEM_REQUIRED, id_columns=("Source item no.",))
    records = [r for r in records if r.get("Source item no.")]
    return ItemMaster(
        path=path,
        sheet=sheet,
        records=tuple(records),
        _by_item_no=_index(records, lambda r: normalize_id(r.get("Source item no."))),
    )


def _stat(path: str | Path, label: str) -> tuple[str, int]:
    resolved = Path(path).expanduser().resolve()
    if not resolved.is_file():
        raise MasterDataError(f"{label} file not found: {resolved}")
    return str(resolved), resolved.stat().st_mtime_ns


def _index(records: Iterable[Record], key_fn) -> dict[str, list[Record]]:
    index: dict[str, list[Record]] = {}
    for record in records:
        key = key_fn(record)
        if key:
            index.setdefault(key, []).append(record)
    return index


def _header_key(value: Any) -> str:
    return " ".join(str(value).split()).casefold()


def _read_table(
    path: str, sheet_name: str, required: tuple[str, ...], id_columns: tuple[str, ...]
) -> tuple[str, list[Record]]:
    """Read the master sheet, detecting the header row below any title rows."""
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise MasterDataError(f"Cannot open master file {path}: {type(exc).__name__}: {exc}") from exc
    try:
        if sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        elif workbook.worksheets:
            sheet = workbook.worksheets[0]
        else:
            raise MasterDataError(f"Master file {path} has no worksheets")
        title = sheet.title
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    required_keys = {_header_key(c): c for c in required}
    header_index, best_missing = None, list(required)
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        present = {_header_key(c) for c in row if c is not None}
        missing = [name for key, name in required_keys.items() if key not in present]
        if not missing:
            header_index = index
            break
        if len(missing) < len(best_missing):
            best_missing = missing
    if header_index is None:
        raise MasterDataError(
            f"{Path(path).name} (sheet '{title}'): required column(s) {best_missing} not found "
            f"in the first {HEADER_SCAN_ROWS} rows"
        )

    # Map each header cell to its canonical column name (as spelled in `required` when it is one).
    header: list[Optional[str]] = []
    for cell in rows[header_index]:
        if cell is None or str(cell).strip() == "":
            header.append(None)
        else:
            header.append(required_keys.get(_header_key(cell), " ".join(str(cell).split())))

    records: list[Record] = []
    for row in rows[header_index + 1 :]:
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        record = {name: _clean(value) for name, value in zip(header, row) if name}
        for column in id_columns:
            record[column] = normalize_id(record.get(column))
        records.append(record)
    return title, records


def _clean(value: Any) -> Any:
    """Make cell values JSON-serialisable."""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip()
    return value
