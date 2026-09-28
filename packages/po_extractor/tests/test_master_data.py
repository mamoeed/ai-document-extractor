import os
import time

import openpyxl
import pytest

from po_extractor.errors import MasterDataError
from po_extractor.validate.master_data import load_customer_master, load_item_master

from .conftest import CUSTOMER_MASTER, ITEM_MASTER


def test_customer_master_detects_header_below_title_rows():
    master = load_customer_master(CUSTOMER_MASTER)
    assert master.sheet == "Customer Master"
    assert len(master.records) == 10
    first = master.records[0]
    assert first["Customer no."] == "48326"
    assert first["Legal name"] == "ClearLine Hygiene GmbH"
    assert first["Active"] == "Yes"


def test_customer_master_ignores_test_scenarios_sheet():
    master = load_customer_master(CUSTOMER_MASTER)
    assert not any(str(r.get("Customer master ID", "")).startswith("SC-") for r in master.records)
    assert all(r["Customer master ID"].startswith("CM-") for r in master.records)


def test_customer_number_48326_has_two_rows():
    names = {r["Legal name"] for r in load_customer_master(CUSTOMER_MASTER).by_number("48326")}
    assert names == {"ClearLine Hygiene GmbH", "Northstar Commercial Kitchens Ltd."}


def test_item_master_loads():
    master = load_item_master(ITEM_MASTER)
    assert master.sheet == "Item Master"
    assert len(master.records) == 34
    record = master.by_item_no("7842136")[0]
    assert record["Base UoM"] == "pc"
    assert record["Net price EUR"] == pytest.approx(223.6)
    assert master.by_item_no(7842136.0)  # numeric lookups are normalized


def test_cache_returns_same_object_until_file_changes(tmp_path):
    path = tmp_path / "customers.xlsx"
    _write_customer_file(path, [("1", "A GmbH", "Yes")])
    first = load_customer_master(path)
    assert load_customer_master(path) is first

    _write_customer_file(path, [("1", "A GmbH", "Yes"), ("2", "B GmbH", "Yes")])
    future = time.time() + 5
    os.utime(path, (future, future))
    reloaded = load_customer_master(path)
    assert reloaded is not first
    assert len(reloaded.records) == 2


def test_header_detection_with_title_rows_and_first_sheet_fallback(tmp_path):
    path = tmp_path / "items.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Whatever"
    sheet.append(["My item list"])
    sheet.append([])
    sheet.append(["exported 2026"])
    sheet.append(["Source item no.", "Description", "Base UoM", "Net price EUR", "Active"])
    sheet.append([123.0, "Thing", "pc", 1.5, "Yes"])
    workbook.save(path)

    master = load_item_master(path)
    assert master.sheet == "Whatever"
    assert master.by_item_no("123")[0]["Description"] == "Thing"


def test_missing_required_column_raises(tmp_path):
    path = tmp_path / "customers.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Customer Master"
    workbook.active.append(["Customer no.", "Legal name"])  # no 'Active'
    workbook.active.append(["1", "A GmbH"])
    workbook.save(path)
    with pytest.raises(MasterDataError, match="Active"):
        load_customer_master(path)


def test_missing_file_raises(tmp_path):
    with pytest.raises(MasterDataError, match="not found"):
        load_item_master(tmp_path / "nope.xlsx")


def _write_customer_file(path, rows):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Customer Master"
    sheet.append(["Title row"])
    sheet.append(["Customer no.", "Legal name", "Active"])
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
